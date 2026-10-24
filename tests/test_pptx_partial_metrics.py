"""PPT000 when part of a shape's text had no font metrics.

A run without metrics was left out of the layout, and the shape counted as
measured because another run was; a fit verdict on the rest passed silently.
"""
from __future__ import annotations

from pathlib import Path

import pytest
from lxml import etree

from ooxml_integrity import Severity, check_pptx
from ooxml_integrity.coverage import CoverageStatus, pptx_coverage
from ooxml_integrity.fonts import EMU_PER_POINT, Metrics, ResolvedFace
from ooxml_integrity.pptx_layout import Paragraph, Run, Shape, layout_shape

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"


@pytest.fixture
def metrics(monkeypatch):
    """Fixed widths for every font but "Missing", which has none."""
    m = Metrics(ResolvedFace("Test", Path("unused.ttf"), "Test", "exact"),
                1000, 800, 200, 0, {cp: 100 for cp in range(32, 127)}, 100, {})
    monkeypatch.setattr("ooxml_integrity.pptx_layout._metrics_for",
                        lambda run: None if run.font == "Missing" else m)
    monkeypatch.setattr("ooxml_integrity.pptx_checks.measurement_available",
                        lambda: (True, "test metrics"))
    return m


def shape(*runs, height=60):
    return Shape("box", 1, 0, 0, round(400 * EMU_PER_POINT), round(height * EMU_PER_POINT),
                 paragraphs=[Paragraph([Run(t, 18, f, False, False) for t, f in runs])],
                 insets=(0, 0, 0, 0))


def deck(tmp_path, *runs, height=60):
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation()
    box = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(
        Pt(36), Pt(36), Pt(400), Pt(height))
    old = box.text_frame._txBody
    body = "".join(f'<a:r><a:rPr lang="en-US" sz="1800"><a:latin typeface="{f}"/></a:rPr>'
                   f"<a:t>{t}</a:t></a:r>" for t, f in runs)
    old.getparent().replace(old, etree.fromstring(
        f'<p:txBody xmlns:p="{P}" xmlns:a="{A}"><a:bodyPr lIns="0" tIns="0" rIns="0" '
        f'bIns="0"><a:noAutofit/></a:bodyPr><a:lstStyle/><a:p>{body}</a:p></p:txBody>'))
    path = tmp_path / "deck.pptx"
    prs.save(path)
    return path


def test_the_layout_names_the_fonts_it_left_out(metrics):
    result = layout_shape(shape(("Fits ", "Test"), ("left out", "Missing")))
    assert result.measured and result.unmeasured_fonts == ["Missing"]
    assert layout_shape(shape(("Fits", "Test"))).unmeasured_fonts == []


def test_part_of_the_text_unmeasured_is_ppt000(metrics, tmp_path):
    path = deck(tmp_path, ("Fits ", "Test"), ("a lot more text " * 40, "Missing"))
    found = [(f.code, f.severity) for f in check_pptx(path) if f.severity >= Severity.WARN]
    assert ("PPT000", Severity.ERROR) in found          # before: nothing at all
    message = next(f.message for f in check_pptx(path) if f.code == "PPT000")
    assert message == ("no font metrics were available for 'Missing', so part of this "
                       "shape's text was not measured and its fit was checked without it")
    report = pptx_coverage(path, check_pptx(path))
    items = {i.id: i.status for i in report.items}
    assert items["pptx.font-metrics"] is CoverageStatus.SKIPPED
    assert items["pptx.text-overflow"] is CoverageStatus.SKIPPED


def test_text_without_any_metrics_keeps_its_message(metrics, tmp_path):
    path = deck(tmp_path, ("only this", "Missing"))
    [message] = [f.message for f in check_pptx(path) if f.code == "PPT000"]
    assert message.startswith("no font metrics were available for any run")


def test_a_run_of_spaces_without_metrics_is_not_text(metrics):
    assert layout_shape(shape(("Fits", "Test"), ("   ", "Missing"))).unmeasured_fonts == []


def test_no_fonts_at_all_is_reported_once_by_the_machine_check(metrics, tmp_path, monkeypatch):
    monkeypatch.setattr("ooxml_integrity.pptx_checks.measurement_available",
                        lambda: (False, "no fonts"))
    path = deck(tmp_path, ("Fits ", "Test"), ("left out", "Missing"))
    assert [f.where for f in check_pptx(path) if f.code == "PPT000"] == [""]


def test_every_run_measured_is_no_ppt000(metrics, tmp_path):
    path = deck(tmp_path, ("Fits ", "Test"), ("too", "Test"))
    assert [f for f in check_pptx(path) if f.code == "PPT000"] == []
