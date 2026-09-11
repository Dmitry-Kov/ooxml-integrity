"""Grow-shape autofit is an explicit coverage gap, not a clean verdict."""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Pt

from ooxml_integrity import check_pptx
from ooxml_integrity.cli import EXIT_FINDINGS, EXIT_OK, main
from ooxml_integrity.coverage import pptx_coverage
from ooxml_integrity.fonts import Metrics, ResolvedFace
from ooxml_integrity.pptx_layout import layout_shape, read_deck


GROW = "pptx.autofit-grow-shape"
OVERFLOW = "pptx.text-overflow"


@pytest.fixture
def metrics(monkeypatch):
    # Ten points per character at 20pt, independent of the host font set.
    face = ResolvedFace("Test", Path("unused.ttf"), "Test", "exact")
    measured = Metrics(face, 1000, 800, 200, 0,
                       {cp: 500 for cp in range(32, 127)}, 500, {})
    monkeypatch.setattr("ooxml_integrity.pptx_layout._metrics_for", lambda run: measured)
    monkeypatch.setattr("ooxml_integrity.fonts.resolve_face", lambda *args: face)
    monkeypatch.setattr("ooxml_integrity.coverage.resolve_face", lambda *args: face)
    monkeypatch.setattr("ooxml_integrity.pptx_checks.measurement_available",
                        lambda: (True, "test metrics"))


def make_deck(tmp_path, modes, *, overlapping=False):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index, mode in enumerate(modes):
        box = slide.shapes.add_textbox(
            Pt(-10 if overlapping else 20 + index * 80), Pt(20), Pt(50), Pt(20),
        )
        box.name = f"{mode}-{index}"
        frame = box.text_frame
        frame.auto_size = (MSO_AUTO_SIZE.NONE if mode == "fixed"
                           else MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT)
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Pt(0)
        frame.margin_top = frame.margin_bottom = Pt(0)
        if mode != "blank-grow":
            run = frame.paragraphs[0].add_run()
            run.text = "W" * 20
            run.font.name = "Test"
            run.font.size = Pt(20)
    path = tmp_path / "autofit.pptx"
    presentation.save(path)
    return path


def inventory(path, findings):
    return {item.id: item.as_dict() for item in pptx_coverage(path, findings).items}


@pytest.mark.parametrize("modes,overflow_count,overflow_status,grow_count", [
    ([], 0, "not-present", 0),
    (["grow"], 0, "skipped", 1),
    (["grow", "grow"], 0, "skipped", 2),
    (["fixed"], 1, "checked", 0),
    (["grow", "fixed"], 1, "checked", 1),
    (["blank-grow"], 0, "not-present", 0),
])
def test_overflow_counts_only_shapes_receiving_a_verdict(
        tmp_path, metrics, modes, overflow_count, overflow_status, grow_count):
    path = make_deck(tmp_path, modes)
    findings = check_pptx(path)
    items = inventory(path, findings)

    assert items[OVERFLOW]["count"] == overflow_count
    assert items[OVERFLOW]["status"] == overflow_status
    assert items[GROW]["count"] == grow_count
    assert items[GROW]["status"] == ("skipped" if grow_count else "not-present")
    if grow_count:
        assert "spAutoFit" in items[GROW]["reason"]
        assert "not checked" in items[GROW]["reason"]
        assert "excluded" in items[OVERFLOW]["reason"]
    # Identical text/rectangles overflow in fixed mode; grow mode gets no verdict.
    assert [(f.code, f.where) for f in findings] == [
        ("PPT001", f"slide1/fixed-{index}")
        for index, mode in enumerate(modes) if mode == "fixed"
    ]


@pytest.mark.parametrize("modes", [["grow"], ["grow", "fixed"]])
def test_missing_metrics_do_not_hide_the_autofit_skip(
        tmp_path, metrics, monkeypatch, modes):
    path = make_deck(tmp_path, modes)
    monkeypatch.setattr("ooxml_integrity.pptx_layout._metrics_for", lambda run: None)
    monkeypatch.setattr("ooxml_integrity.pptx_checks.measurement_available",
                        lambda: (False, "test fonts unavailable"))
    findings = check_pptx(path)
    items = inventory(path, findings)

    assert [f.code for f in findings] == ["PPT000"]
    assert items["pptx.font-metrics"]["status"] == "skipped"
    assert items[OVERFLOW]["status"] == "skipped"
    assert items[OVERFLOW]["count"] == modes.count("fixed")
    assert items[GROW]["status"] == "skipped"
    assert items[GROW]["count"] == 1
    assert "spAutoFit" in items[GROW]["reason"]
    assert "not checked" in items[GROW]["reason"]


@pytest.mark.parametrize("available", [True, False])
def test_grow_shapes_still_receive_geometry_findings(
        tmp_path, metrics, monkeypatch, available):
    path = make_deck(tmp_path, ["grow", "grow"], overlapping=True)
    if not available:
        monkeypatch.setattr("ooxml_integrity.pptx_layout._metrics_for", lambda run: None)
        monkeypatch.setattr("ooxml_integrity.pptx_checks.measurement_available",
                            lambda: (False, "test fonts unavailable"))
    findings = check_pptx(path)
    assert [f.code for f in findings if f.code != "PPT000"] == [
        "PPT004", "PPT004", "PPT006",
    ]
    items = inventory(path, findings)
    for identifier in ("pptx.off-slide-geometry", "pptx.text-shape-overlap"):
        assert items[identifier]["status"] == "checked"
        assert items[identifier]["count"] == 2
    assert items[OVERFLOW]["count"] == 0
    assert items[GROW]["status"] == "skipped"


def test_grow_shapes_do_not_affect_overflow_layout_confidence(
        tmp_path, metrics, monkeypatch):
    path = make_deck(tmp_path, ["grow", "fixed"])
    measured = []

    def record(shape):
        measured.append(shape.name)
        return layout_shape(shape)

    monkeypatch.setattr("ooxml_integrity.coverage.layout_shape", record)
    items = inventory(path, check_pptx(path))
    assert measured == ["fixed-1"]
    assert items[OVERFLOW]["status"] == "checked"


@pytest.mark.parametrize("json_output", [False, True], ids=["human", "json"])
@pytest.mark.parametrize("modes", [[], ["grow"], ["grow", "fixed"]],
                         ids=["empty", "grow-only", "mixed"])
def test_cli_exposes_the_skip_without_changing_exit_code(
        tmp_path, metrics, capsys, json_output, modes):
    path = make_deck(tmp_path, modes)
    # This shape exceeds its stored rectangle; that is not new renderer evidence.
    for shape in read_deck(path).shapes:
        assert layout_shape(shape).vertical_overflow_ratio > 2
    exit_code = EXIT_FINDINGS if "fixed" in modes else EXIT_OK
    assert main(["check", str(path)]) == exit_code
    capsys.readouterr()
    args = ["check", str(path), "--coverage"]
    if json_output:
        args.append("--json")
    assert main(args) == exit_code
    output = capsys.readouterr().out

    if json_output:
        file_report = json.loads(output)["files"][0]
        assert [f["code"] for f in file_report["findings"]] == [
            "PPT001" for mode in modes if mode == "fixed"
        ]
        report = file_report["coverage"]
        assert report["schema_version"] == 1
        items = {item["id"]: item for item in report["items"]}
        assert items[GROW]["status"] == ("skipped" if modes else "not-present")
        assert items[GROW]["count"] == modes.count("grow")
        assert items[OVERFLOW]["status"] == (
            "checked" if "fixed" in modes else "skipped" if modes else "not-present"
        )
        assert items[OVERFLOW]["count"] == modes.count("fixed")
    else:
        if exit_code == EXIT_OK:
            assert "no findings in checked surfaces" in output
        if modes:
            assert f"[skipped] {GROW}" in output
            assert "spAutoFit" in output
            assert "not checked" in output
        else:
            assert GROW not in output  # Concise output omits absent surfaces.
