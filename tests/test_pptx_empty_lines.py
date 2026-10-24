"""Empty paragraphs and empty lines take height, and the anchor decides whether it shows.

Before, an empty paragraph added nothing, an empty line between breaks or after a
trailing break was always 18pt, and blank rows below the last line of text counted
even where they move no glyph. The fixture's renders are in
docs/calibration/pptx-empty-lines/evidence.json, each with the tolerance its
measurements need.
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pytest
from lxml import etree

from conftest import ROOT, needs_zlib_bytes, read_part, repack
from ooxml_integrity import Severity, check_pptx
from ooxml_integrity.fonts import EMU_PER_POINT, Metrics, ResolvedFace
from ooxml_integrity.pptx_checks import check_text_overflow
from ooxml_integrity.pptx_layout import Deck, Paragraph, Run, Shape, layout_shape, read_deck
from research.pptx_empty_lines import ANCHORS, DELTAS, build

EVIDENCE = json.loads((ROOT / "docs/calibration/pptx-empty-lines/evidence.json").read_text())
FIXTURE = ROOT / EVIDENCE["fixture"]


@pytest.fixture
def metrics(monkeypatch):
    # Heights do not depend on advance widths; keep lines from wrapping anyway.
    m = Metrics(ResolvedFace("Test", Path("unused.ttf"), "Test", "exact"),
                1000, 800, 200, 0, {cp: 100 for cp in range(32, 127)}, 100, {})
    monkeypatch.setattr("ooxml_integrity.pptx_layout._metrics_for", lambda run: m)
    monkeypatch.setattr("ooxml_integrity.pptx_checks.measurement_available",
                        lambda: (True, "test metrics"))
    return m


def text(t="Abc", size=18):
    return Paragraph([Run(t, size, "Test", False, False)])


def empty(end=18, **kwargs):
    return Paragraph([], end_size_pt=end, **kwargs)


def shape(*paragraphs, height=500, anchor="t", **kwargs):
    return Shape("box", 1, 0, 0, round(500 * EMU_PER_POINT), round(height * EMU_PER_POINT),
                 paragraphs=list(paragraphs), insets=(0, 0, 0, 0), anchor=anchor, **kwargs)


def height(*paragraphs, **kwargs):
    return layout_shape(shape(*paragraphs, **kwargs)).text_height_pt


def overflow(s):
    return [f.code for f in check_text_overflow(Deck(Path("t.pptx"), 9144000, 6858000, [s], {}))
            if f.severity >= Severity.WARN]


# ------------------------------------------------------------ the model
def test_an_empty_paragraph_is_as_tall_as_its_mark(metrics):
    base = height(text(), text("Xyz"))
    assert height(text(), empty(36), text("Xyz")) - base == pytest.approx(43.2)
    assert height(text(), empty(10), text("Xyz")) - base == pytest.approx(12.0)


def test_an_empty_line_is_as_tall_as_the_break_or_mark_that_ends_it(metrics):
    def broken(*runs, end=18):
        return Paragraph([Run(t, s, "Test", False, False) for t, s in runs], end_size_pt=end)
    base = height(broken(("Abc", 18), ("\n", 36), ("Xyz", 18)))
    between = height(broken(("Abc", 18), ("\n", 36), ("\n", 10), ("Xyz", 18)))
    assert between - base == pytest.approx(12.0)
    trailing = height(broken(("Abc", 18), ("\n", 18), end=36), text("Xyz"))
    assert trailing - height(text(), text("Xyz")) == pytest.approx(43.2)


def test_an_empty_paragraph_keeps_its_spacing(metrics):
    base = height(text(), text("Xyz"))
    spaced = empty(18, space_before_pt=12, space_after_pt=6)
    assert height(text(), spaced, text("Xyz")) - base == pytest.approx(21.6 + 12 + 6)


def test_stored_shrink_scales_empty_lines_too(metrics):
    s = shape(text(), empty(40), text("Xyz"), autofit="normAutofit", font_scale=0.5)
    assert layout_shape(s).text_height_pt == pytest.approx(0.5 * (21.6 + 48.0 + 21.6))


@pytest.mark.parametrize("anchor, expected", [
    ("t", 43.2 + 86.4),     # blank rows above push the text down
    ("b", 43.2),            # bottom-anchored text ends at the bottom regardless
    ("ctr", 43.2 + 86.4),   # centred text moves by half the difference, up or down
])
def test_leading_blank_rows_count_by_anchor(metrics, anchor, expected):
    assert height(empty(36), empty(36), text(), text("Xyz"), anchor=anchor) == pytest.approx(expected)


@pytest.mark.parametrize("anchor, expected", [
    ("t", 43.2), ("b", 43.2 + 86.4), ("ctr", 43.2 + 86.4),
])
def test_trailing_blank_rows_count_by_anchor(metrics, anchor, expected):
    assert height(text(), text("Xyz"), empty(36), empty(36), anchor=anchor) == pytest.approx(expected)


def test_centred_text_with_blank_rows_on_both_sides_moves_by_their_difference(metrics):
    assert height(empty(36), text(), empty(36), anchor="ctr") == pytest.approx(21.6)
    assert height(empty(36), empty(36), text(), empty(36), anchor="ctr") == pytest.approx(21.6 + 43.2)


def test_no_space_below_the_last_paragraph(metrics):
    for anchor in ("t", "ctr", "b"):
        last = Paragraph([Run("Xyz", 18, "Test", False, False)], space_after_pt=36)
        assert height(text(), last, anchor=anchor) == pytest.approx(43.2)
        middle = Paragraph([Run("Abc", 18, "Test", False, False)], space_after_pt=36)
        assert height(middle, text("Xyz"), anchor=anchor) == pytest.approx(43.2 + 36)


def test_empty_paragraphs_between_text_overflow_a_box(metrics):
    s = shape(text(), empty(), empty(), empty(), text("Xyz"), height=50)
    assert overflow(s) == ["PPT001"]          # before: 43.2pt, inside the box


def test_trailing_empty_paragraphs_in_a_top_anchored_box_are_clean(metrics):
    # One line of text and four empty paragraphs below it: nothing is drawn
    # outside the box, so nothing is reported.
    s = shape(text(), empty(), empty(), empty(), empty(), height=50)
    assert overflow(s) == []
    assert overflow(shape(text(), empty(), empty(), empty(), empty(), height=50,
                          anchor="b")) == ["PPT001"]


def test_trailing_breaks_in_a_top_anchored_box_are_clean(metrics):
    # The aascu.org title in Apache POI's corpus: three lines at 90% spacing,
    # then three breaks.
    runs = [Run(t, 26, "Test", False, False)
            for t in ("A", "\n", "B", "\n", "C", "\n", "", "\n", "", "\n")]
    s = shape(Paragraph(runs, line_spacing=0.9, end_size_pt=36), height=88.8)
    assert overflow(s) == []                  # before: 82% over


def test_unmeasured_text_is_still_unmeasured_next_to_empty_paragraphs(monkeypatch):
    monkeypatch.setattr("ooxml_integrity.pptx_layout._metrics_for", lambda run: None)
    result = layout_shape(shape(text(), empty(), empty()))
    assert not result.measured


# ---------------------------------------------------------- reading files
def test_marks_and_breaks_are_read_with_their_inherited_sizes():
    deck = {s.name: s for s in read_deck(FIXTURE).shapes}
    assert [p.end_size_pt for p in deck["empty36"].paragraphs] == [18, 36, 18]
    assert deck["empty_inherited"].paragraphs[1].end_size_pt == 18   # presentation default
    assert [r.size_pt for r in deck["break_break10"].paragraphs[0].runs] == [18, 10, 10, 18]


def test_a_placeholder_inherits_its_anchor(tmp_path):
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])         # title slide
    slide.shapes.title.text = "Title"
    slide.placeholders[1].text = "Subtitle"
    path = tmp_path / "ph.pptx"
    prs.save(path)
    anchors = {s.placeholder_type: s.anchor for s in read_deck(path).shapes}
    # The layout's ctrTitle has no anchor; the master's title placeholder is ctr.
    assert anchors == {"ctrTitle": "ctr", "subTitle": "t"}

    layout = "ppt/slideLayouts/slideLayout1.xml"
    tree = etree.fromstring(read_part(path, layout).encode())
    a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    tree.findall(f".//{a}bodyPr")[1].set("anchor", "b")
    edited = repack(path, tmp_path / "b.pptx", {layout: etree.tostring(tree)})
    anchors = {s.placeholder_type: s.anchor for s in read_deck(edited).shapes}
    assert anchors["subTitle"] == "b"


def test_a_placeholder_inherits_from_the_layout_placeholder_with_its_idx(tmp_path):
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[3])         # two content
    for ph in slide.placeholders:
        ph.text = f"idx {ph.placeholder_format.idx}"
    path = tmp_path / "two.pptx"
    prs.save(path)
    # Both content placeholders are of type body; the left one comes first.
    layout = "ppt/slideLayouts/slideLayout4.xml"
    tree = etree.fromstring(read_part(path, layout).encode())
    p = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    for sp in tree.iter(f"{p}sp"):
        idx = sp.find(f"{p}nvSpPr/{p}nvPr/{p}ph").get("idx")
        if idx in ("1", "2"):
            sp.find(f"{p}txBody/{a}bodyPr").set("anchor", {"1": "t", "2": "b"}[idx])
    edited = repack(path, tmp_path / "anchored.pptx", {layout: etree.tostring(tree)})
    anchors = {s.paragraphs[0].runs[0].text: s.anchor for s in read_deck(edited).shapes}
    assert anchors["idx 1"] == "t" and anchors["idx 2"] == "b"


# ------------------------------------------------------- renderer evidence
@needs_zlib_bytes
def test_the_fixture_rebuilds_byte_for_byte(tmp_path):
    out = tmp_path / "deck.pptx"
    build(out)
    digest = hashlib.sha256(out.read_bytes()).hexdigest()
    assert digest == EVIDENCE["fixture_sha256"]
    assert hashlib.sha256(FIXTURE.read_bytes()).hexdigest() == digest


@pytest.mark.parametrize("renderer", sorted(EVIDENCE["renderers"]))
def test_model_heights_match_each_renderer(renderer):
    probes = EVIDENCE["renderers"][renderer]["probes"]
    tolerance = EVIDENCE["renderers"][renderer]["tolerance_pt"]
    model = {s.name: layout_shape(s) for s in read_deck(FIXTURE).shapes}
    for probe, against in DELTAS:
        drawn = probes[probe]["x_bottom_pt"] - probes[against]["x_bottom_pt"]
        predicted = model[probe].text_height_pt - model[against].text_height_pt
        assert predicted == pytest.approx(drawn, abs=tolerance), (renderer, probe)


@pytest.mark.parametrize("renderer", sorted(EVIDENCE["renderers"]))
def test_overflow_verdicts_match_what_each_renderer_drew(renderer):
    probes = EVIDENCE["renderers"][renderer]["probes"]
    edge = max(0.5, EVIDENCE["renderers"][renderer]["tolerance_pt"])
    # Glyph boxes sit inside the line boxes; t_text puts the lines at 0..43.2pt.
    top_inset = probes["t_text"]["a_top_pt"]
    bottom_inset = probes["t_text"]["x_bottom_pt"] - 43.2
    flagged = {f.where.split("/", 1)[1] for f in check_pptx(FIXTURE)
               if f.code == "PPT001"}
    for name, (_, box, _) in ANCHORS.items():
        drawn = probes[name]
        outside = (drawn["a_top_pt"] - top_inset < -edge
                   or drawn["x_bottom_pt"] - bottom_inset > box + edge)
        assert (name in flagged) == outside, (renderer, name)
    assert flagged <= set(ANCHORS)
