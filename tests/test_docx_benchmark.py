"""Fault injection checks the benchmark, not the quality of an editor.

Synthetic mutations here are never counted as real adapter attempts.
"""
import ast
import json
import shutil

import pytest
from lxml import etree as E

from research import docx_benchmark as bench
from research import revision_evidence as rev


def source(profile="basic"):
    return rev.BASE / "sources" / f"{profile}.docx"


def untracked_parts(profile="basic"):
    parts = rev.read(source(profile))
    root = E.fromstring(parts[rev.MAIN])
    target = next(t for t in root.iter(rev.W + "t") if bench.OLD in (t.text or ""))
    target.text = target.text.replace(bench.OLD, bench.NEW)
    parts[rev.MAIN] = rev.xml(root)
    return parts


def output(tmp_path, parts):
    path = tmp_path / "output.docx"
    rev.write(parts, path)
    return path


def test_protocol_is_frozen_and_every_input_has_the_declared_plain_target():
    protocol = bench.frozen_protocol()
    assert set(protocol["sources"]) == set(rev.PROFILES)
    assert protocol["agent"]["status"] == "not_run"
    for item in protocol["sources"].values():
        root = E.fromstring(rev.read(bench.ROOT / item["path"])[rev.MAIN])
        target = [t for t in root.iter(rev.W + "t") if bench.OLD in (t.text or "")]
        assert len(target) == 1
        run = target[0].getparent()
        assert run.tag == rev.W + "r" and len(run) == 1
        assert run.getparent() == root.findall("w:body/w:p", rev.NS)[2]


@pytest.mark.parametrize("profile", rev.PROFILES)
def test_unchanged_output_is_not_a_completed_edit_but_has_no_collateral_loss(profile):
    result = bench.inspect_pair(source(profile), source(profile), "edit", "python-docx-run-v1")
    assert result["requested_change_completed"] is False
    assert result["protected_content_preserved"] is True
    assert result["actionable_findings"] == []
    assert result["checker_detection"] == "no_independent_defect"
    assert bench.assess(source(profile), source(profile), "save", "python-docx-run-v1")["requested_change_completed"]


@pytest.mark.parametrize("profile", rev.PROFILES)
def test_exact_untracked_edit_preserves_content(profile, tmp_path):
    result = bench.assess(source(profile), output(tmp_path, untracked_parts(profile)), "edit", "python-docx-run-v1")
    assert result["requested_change_completed"]
    assert result["protected_content_preserved"], result["violations"]


@pytest.mark.parametrize("defect,profile", [("replace-unrelated-insertion", "basic"), ("unwrap-note-insertion", "notes")])
def test_existing_checker_misses_stay_losses_even_when_task_is_completed(defect, profile, tmp_path):
    parts = rev.mutate(untracked_parts(profile), defect)
    result = bench.inspect_pair(source(profile), output(tmp_path, parts), "edit", "python-docx-run-v1")
    assert result["requested_change_completed"]
    assert not result["protected_content_preserved"]
    assert result["violations"]["unrelated_revision_losses"]
    assert result["actionable_findings"] == []


@pytest.mark.parametrize("defect", ["anchor", "style", "relationship", "asset", "cell"])
def test_protected_package_rejects_count_neutral_changes(defect, tmp_path):
    parts = untracked_parts()
    root = E.fromstring(parts[rev.MAIN])
    if defect == "anchor":
        anchor = root.find(".//w:commentRangeStart", rev.NS)
        anchor.getparent().remove(anchor)
        root.findall("w:body/w:p", rev.NS)[1].insert(0, anchor)
    elif defect == "style":
        root.find(".//w:pStyle", rev.NS).set(rev.W + "val", "Normal")
    elif defect == "relationship":
        rels = E.fromstring(parts["word/_rels/document.xml.rels"])
        rel = next(n for n in rels if n.get("TargetMode") == "External")
        rel.set("Target", "https://example.org/different")
        parts["word/_rels/document.xml.rels"] = rev.xml(rels)
    elif defect == "asset":
        binary = next(p for p in parts if p.startswith("word/media/"))
        parts[binary] += b"changed"
    else:
        cells = root.findall(".//w:tc", rev.NS)
        row = cells[0].getparent()
        row.remove(cells[0])
        row.append(cells[0])
    parts[rev.MAIN] = rev.xml(root)
    result = bench.assess(source(), output(tmp_path, parts), "edit", "python-docx-run-v1")
    assert result["requested_change_completed"]
    assert not result["protected_content_preserved"]
    assert result["violations"]["protected_package_changes"]


def tracked_parts(formatted=False, utc_extension=False):
    parts = rev.read(source())
    root = E.fromstring(parts[rev.MAIN])
    paragraph = root.findall("w:body/w:p", rev.NS)[2]
    run = paragraph[0]
    prefix, suffix = run[0].text.split(bench.OLD)
    paragraph.remove(run)
    elements = [E.fromstring(f'<w:r xmlns:w="{rev.W[1:-1]}"><w:t>{prefix}</w:t></w:r>'),
                E.fromstring(f'<w:del xmlns:w="{rev.W[1:-1]}" w:id="900" w:author="Evidence Editor">'
                             '<w:r><w:delText>EDITBEFORE</w:delText></w:r></w:del>'),
                E.fromstring(f'<w:ins xmlns:w="{rev.W[1:-1]}" w:id="901" w:author="Evidence Editor">'
                             '<w:r><w:t>EDITAFTER</w:t></w:r></w:ins>'),
                E.fromstring(f'<w:r xmlns:w="{rev.W[1:-1]}"><w:t>{suffix}</w:t></w:r>')]
    if formatted:
        prop = E.Element(rev.W + "rPr")
        prop.append(E.Element(rev.W + "b"))
        elements[2][0].insert(0, prop)
    if utc_extension:
        for node in elements[1:3]:
            node.set(rev.W + "date", rev.DATE)
            node.set("{http://schemas.microsoft.com/office/word/2023/wordml/word16du}dateUtc", rev.DATE)
    for i, element in enumerate(elements):
        paragraph.insert(i, element)
    parts[rev.MAIN] = rev.xml(root)
    return parts


@pytest.mark.parametrize("formatted", [False, True])
@pytest.mark.parametrize("utc_extension", [False, True])
def test_tracked_marker_normalization_does_not_hide_new_formatting(formatted, utc_extension, tmp_path):
    result = bench.assess(source(), output(tmp_path, tracked_parts(formatted, utc_extension)), "edit", "adeu-sdk-v1")
    assert result["requested_change_completed"]
    assert result["protected_content_preserved"] is (not formatted)


def test_adeu_contract_requires_new_tracking_even_with_correct_current_text(tmp_path):
    result = bench.assess(source(), output(tmp_path, untracked_parts()), "edit", "adeu-sdk-v1")
    assert not result["requested_change_completed"]
    assert result["protected_content_preserved"]


def test_adapter_refuses_missing_or_split_target_without_writing_output(tmp_path):
    from docx import Document
    path = tmp_path / "split.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("EDIT")
    paragraph.add_run("BEFORE")
    document.save(path)
    destination = tmp_path / "absent.docx"
    with pytest.raises(bench.Unsupported):
        bench.edit_python_docx(path, destination, "edit")
    assert not destination.exists()


def test_receipts_refuse_overwrite(tmp_path):
    path = tmp_path / "receipt.json"
    bench.save_json(path, {"original": True})
    with pytest.raises(FileExistsError):
        bench.save_json(path, {"replacement": True})
    assert json.loads(path.read_text()) == {"original": True}


@pytest.mark.frozen_checker
def test_captured_receipts_and_evaluations_replay_when_present():
    captures = list((bench.BASE / "captures").glob("*/capture.json"))
    if not captures:
        pytest.skip("Fresh captures have not been created yet")
    assert len(captures) == 4
    for path in captures:
        actual = bench.evaluate(path.parent)
        expected = json.loads((path.parent / "evaluation.reviewed.json").read_text())
        assert actual == expected
        assert all(c["overall_success"] for c in actual["cases"])


def test_initial_and_reviewed_capture_adapters_have_identical_code():
    def adapters(path):
        tree = ast.parse(path.read_text())
        return {node.name: ast.dump(node) for node in tree.body
                if isinstance(node, ast.FunctionDef) and node.name in ("edit_adeu", "edit_python_docx")}
    current = adapters(bench.ROOT / "research/docx_benchmark.py")
    for path in (bench.BASE / "captures").glob("*/capture-harness.py"):
        assert adapters(path) == current


@pytest.mark.frozen_checker
def test_verification_rejects_tampered_output(tmp_path):
    original = bench.BASE / "captures/python-docx-1"
    if not original.exists():
        pytest.skip("No capture yet")
    directory = tmp_path / "tampered"
    shutil.copytree(original, directory)
    path = directory / "basic-edit.docx"
    path.write_bytes(path.read_bytes() + b"tampered")
    with pytest.raises(ValueError, match="output hash drift"):
        bench.evaluate(directory)


def test_new_seeded_anchor_gap_is_separate_and_reproducible():
    path = bench.BASE / "diagnostics/seeded-comment-anchor-move.docx"
    receipt = json.loads(path.with_suffix(".json").read_text())
    assert receipt["kind"] == "seeded_diagnostic_not_editor_output"
    assert receipt["output_sha256"] == bench.digest(path)
    parts = rev.read(source())
    root = E.fromstring(parts[rev.MAIN])
    anchor = root.find(".//w:commentRangeStart", rev.NS)
    anchor.getparent().remove(anchor)
    root.findall("w:body/w:p", rev.NS)[1].insert(0, anchor)
    parts[rev.MAIN] = rev.xml(root)
    assert parts == rev.read(path)
    result = bench.inspect_pair(source(), path, "save", "python-docx-run-v1")
    assert json.loads(bench.json_bytes(result)) == receipt["assessment"]
    assert result["requested_change_completed"]
    assert not result["protected_content_preserved"]
    assert result["revision_oracle"]["intent_correct"]  # Original oracle's positional blind spot.
    assert result["actionable_findings"] == []
