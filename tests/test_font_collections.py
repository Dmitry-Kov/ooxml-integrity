"""Collection selection must survive discovery, grading, coverage and metrics."""
from __future__ import annotations

import hashlib
import json
import subprocess
from dataclasses import replace

import pytest
from fontTools.fontBuilder import FontBuilder
from fontTools.pens.t2CharStringPen import T2CharStringPen
from fontTools.pens.ttGlyphPen import TTGlyphPen
from fontTools.ttLib import TTCollection, TTFont
from lxml import etree
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Pt

from conftest import read_part, repack
from ooxml_integrity import Severity, check_pptx
from ooxml_integrity import fonts
from ooxml_integrity.cli import EXIT_FINDINGS, EXIT_OK, main
from ooxml_integrity.doctor import build_report
from ooxml_integrity.pptx_layout import A, layout_shape, read_deck


def make_font(family, style, width, *, cff=False, latin=True):
    """Original empty outlines with deliberately different advance tables."""
    fb = FontBuilder(1000, isTTF=not cff)
    cmap = {cp: f"g{cp}" for cp in (range(32, 127) if latin else [0x4E00])}
    order = [".notdef", *cmap.values()]
    fb.setupGlyphOrder(order)
    fb.setupCharacterMap(cmap)
    if cff:
        fb.setupCFF(family.replace(" ", "") + style.replace(" ", ""), {},
                    {g: T2CharStringPen(width, None).getCharString() for g in order}, {})
    else:
        fb.setupGlyf({g: TTGlyphPen(None).glyph() for g in order})
    fb.setupHorizontalMetrics({g: (width, 0) for g in order})
    fb.setupHorizontalHeader(ascent=700 + width // 10, descent=-200)
    fb.setupNameTable({"familyName": family, "styleName": style,
                       "typographicFamily": family, "typographicSubfamily": style,
                       "fullName": family + " " + style,
                       "psName": family.replace(" ", "") + style.replace(" ", "")})
    fb.setupOS2(sTypoAscender=700 + width // 10, sTypoDescender=-200,
                usWinAscent=900, usWinDescent=200)
    fb.setupPost()
    return fb.font


@pytest.fixture(autouse=True)
def clear_font_caches():
    cached = (fonts._fc_match, fonts._index_font_dirs, fonts._has_latin_coverage,
              fonts.resolve_face, fonts.load_metrics)
    def clear():
        for fn in cached:
            fn.cache_clear()
    clear()
    yield
    clear()


@pytest.fixture(params=["ttc", "otc"])
def collection(tmp_path, monkeypatch, request):
    cff = request.param == "otc"
    coll = TTCollection()
    coll.fonts = [make_font("A Decoy", "Regular", 900, cff=cff, latin=False)]
    coll.fonts += [make_font("Collection Test", style, width, cff=cff)
                   for style, width in [("Regular", 300), ("Bold", 600),
                                        ("Italic", 200), ("Bold Italic", 700),
                                        ("Light", 100)]]
    path = tmp_path / ("faces." + request.param)
    coll.save(path)
    coll.close()
    monkeypatch.setattr(fonts, "FONT_DIRS", (str(tmp_path),))
    monkeypatch.setattr(fonts, "_fc_match", lambda pattern: None)
    return path


@pytest.mark.parametrize("bold,italic,index,width", [
    (False, False, 1, 300), (True, False, 2, 600),
    (False, True, 3, 200), (True, True, 4, 700),
])
def test_directory_selection_and_metric_cache_preserve_each_member(
        collection, bold, italic, index, width):
    m = fonts.load_metrics("Collection Test", bold, italic)
    assert m.face.face_index == index
    assert m.face.path == collection
    assert m.face.match == "exact"
    assert m.text_width_pt("ABCD", 10) == width / 25
    assert m.ascender == 700 + width // 10
    # A subsequent regular lookup must not reuse the styled member's metrics.
    assert fonts.load_metrics("Collection Test").face.face_index == 1
    with TTFont(collection, fontNumber=index) as direct:
        assert m.widths[ord("A")] == direct["hmtx"][direct.getBestCmap()[ord("A")]][0]


@pytest.mark.parametrize("bold,italic,index,text_width,box_width", [
    (False, False, 1, 12, 18), (True, False, 2, 24, 18),
    (False, True, 3, 8, 10), (True, True, 4, 28, 18),
], ids=["regular-control", "bold-overflow", "italic-fits", "bold-italic-overflow"])
def test_inherited_style_selects_the_face_used_for_layout_and_cli(
        collection, tmp_path, capsys, bold, italic, index, text_width, box_width):
    # Real synthetic font tables make style selection observable without any
    # host-font dependency. Italic is narrower here; bold variants are wider.
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Pt(20), Pt(20), Pt(box_width), Pt(30))
    box.name = "inherited-style"
    frame = box.text_frame
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.word_wrap = False
    frame.margin_left = frame.margin_right = Pt(0)
    frame.margin_top = frame.margin_bottom = Pt(0)
    run = frame.paragraphs[0].add_run()
    run.text = "ABCD"
    run.font.name = "Collection Test"
    run.font.size = Pt(10)
    level = etree.SubElement(frame._txBody.find(f"{{{A}}}lstStyle"), f"{{{A}}}lvl1pPr")
    etree.SubElement(level, f"{{{A}}}defRPr", b=str(int(bold)), i=str(int(italic)))
    path = tmp_path / "inherited-style.pptx"
    presentation.save(path)

    shape, = read_deck(path).shapes
    parsed = shape.paragraphs[0].runs[0]
    measured = fonts.load_metrics(parsed.font, parsed.bold, parsed.italic)
    assert measured.face.path == collection
    assert measured.face.face_index == index
    assert measured.face.match == "exact"
    result = layout_shape(shape)
    assert result.measured and result.confident
    assert result.widest_line_pt == text_width
    regular = replace(shape, paragraphs=[replace(
        shape.paragraphs[0], runs=[replace(parsed, bold=False, italic=False)],
    )])
    assert layout_shape(regular).widest_line_pt == 12
    if bold or italic:
        assert (result.horizontal_overflow_pt > 0) != (
            layout_shape(regular).horizontal_overflow_pt > 0
        )
    overflow = text_width > box_width
    assert [(f.code, f.severity) for f in check_pptx(path)] == (
        [("PPT003", Severity.ERROR)] if overflow else []
    )
    assert main(["check", str(path), "--coverage", "--json"]) == (
        EXIT_FINDINGS if overflow else EXIT_OK
    )
    file_report = json.loads(capsys.readouterr().out)["files"][0]
    assert [f["code"] for f in file_report["findings"]] == (["PPT003"] if overflow else [])
    coverage = {item["id"]: item for item in file_report["coverage"]["items"]}
    assert coverage["pptx.font-metrics"]["status"] == "checked"
    assert coverage["pptx.text-overflow"]["status"] == "checked"


def test_latin_coverage_cache_and_last_resort_use_the_selected_member(collection, monkeypatch):
    assert not fonts._has_latin_coverage(str(collection), 0)
    assert fonts._has_latin_coverage(str(collection), 1)
    assert not fonts._has_latin_coverage(str(collection), 0)
    monkeypatch.setattr(fonts, "LAST_RESORT", ())
    found = fonts.resolve_face("Missing Family")
    assert found.face_index == 1 and found.match == "fallback"
    assert found.family == "collection test"


@pytest.mark.parametrize("grade", ["metric", "similar", "fallback", "first-fallback"])
def test_all_substitution_branches_preserve_nonzero_index(collection, monkeypatch, grade):
    # Exercise both explicit substitute searches and the unclassified first hit.
    family = "Alias"
    target = "Collection Test"
    monkeypatch.setattr(fonts, "METRIC_SUBSTITUTES", {"alias": (target,)} if grade == "metric" else {})
    monkeypatch.setattr(fonts, "SIMILAR_SUBSTITUTES", {"alias": (target,)} if grade == "similar" else {})
    monkeypatch.setattr(fonts, "LAST_RESORT", (target,) if grade == "fallback" else ())
    if grade == "first-fallback":
        monkeypatch.setattr(fonts, "_locate", lambda requested, *args:
                            (str(collection), target, 2) if requested == family else None)
    result = fonts.resolve_face(family, bold=True)
    assert result.face_index == 2
    assert result.match == ("fallback" if grade == "first-fallback" else grade)


def test_directory_scan_keeps_shared_stream_open_and_prefers_regular(collection):
    index = fonts._index_font_dirs()
    assert {number for _, number in index.values()} >= {0, 1, 2, 3, 4}
    assert index["collection test"][1] == 1  # later Light must not overwrite Regular


def test_preferred_subfamily_prevents_heavy_from_impersonating_regular(tmp_path, monkeypatch):
    coll = TTCollection()
    regular = make_font("Preferred Family", "Regular", 300)
    heavy = make_font("Preferred Family", "Regular", 800)
    heavy["name"].setName("Heavy", 17, 3, 1, 0x409)
    coll.fonts = [regular, heavy]
    coll.save(tmp_path / "preferred.ttc")
    coll.close()
    monkeypatch.setattr(fonts, "FONT_DIRS", (str(tmp_path),))
    assert fonts._index_font_dirs()["preferred family"][1] == 0


def test_standalone_font_defaults_to_member_zero(tmp_path, monkeypatch):
    path = tmp_path / "one.ttf"
    f = make_font("Standalone Test", "Regular", 400)
    f.save(path)
    f.close()
    monkeypatch.setattr(fonts, "FONT_DIRS", (str(tmp_path),))
    monkeypatch.setattr(fonts, "_fc_match", lambda pattern: None)
    m = fonts.load_metrics("Standalone Test")
    assert m.face.face_index == 0 and m.text_width_pt("AB", 10) == 8
    assert fonts.ResolvedFace("x", path, "x", "exact").face_index == 0


@pytest.mark.parametrize("grade,requested,resolved", [
    ("exact", "Collection Test", "Collection Test"),
    ("metric", "Calibri", "Carlito"), ("similar", "Segoe UI", "DejaVu Sans"),
])
def test_fontconfig_returns_index_and_resolution_keeps_it(tmp_path, monkeypatch, grade, requested, resolved):
    path = tmp_path / "multi.otc"
    path.touch()
    monkeypatch.setattr(fonts.shutil, "which", lambda name: "/test/fc-match")

    def run(args, **kwargs):
        assert "%{index}" in args[2]
        return subprocess.CompletedProcess(args, 0, f"{path}\t{resolved},Alias\t3\n", "")

    monkeypatch.setattr(fonts.subprocess, "run", run)
    face = fonts.resolve_face(requested)
    assert face.face_index == 3 and face.match == grade


@pytest.mark.parametrize("number", ["", "bogus", "-1", "1\t2"])
def test_fontconfig_does_not_default_malformed_indices_to_zero(tmp_path, monkeypatch, number):
    path = tmp_path / "face.ttc"
    path.touch()
    monkeypatch.setattr(fonts.shutil, "which", lambda name: "/test/fc-match")
    monkeypatch.setattr(fonts.subprocess, "run", lambda args, **kwargs:
                        subprocess.CompletedProcess(args, 0, f"{path}\tTest\t{number}", ""))
    assert fonts._fc_match("Test") is None


def test_named_variable_instance_is_not_silently_masked_to_member_zero(tmp_path, monkeypatch):
    path = tmp_path / "variable.ttf"
    path.touch()
    monkeypatch.setattr(fonts.shutil, "which", lambda name: "/test/fc-match")
    monkeypatch.setattr(fonts.subprocess, "run", lambda args, **kwargs:
                        subprocess.CompletedProcess(args, 0, f"{path}\tTest\t65536", ""))
    with pytest.raises(fonts.FontUnavailable, match="named variable-font instance"):
        fonts.resolve_face("Test")


def test_invalid_collection_member_is_not_replaced_with_zero(collection, monkeypatch):
    monkeypatch.setattr(fonts, "resolve_face", lambda *args:
                        fonts.ResolvedFace("Test", collection, "Test", "exact", 99))
    with pytest.raises(Exception, match="[Ff]ont[Nn]umber|font number"):
        fonts.load_metrics("Test")


def test_doctor_exposes_member_index(collection, monkeypatch):
    monkeypatch.setattr("ooxml_integrity.doctor.FONT_PROBES", ("Collection Test",))
    probe = build_report()["fonts"]["probes"][0]
    assert probe["face_index"] == 1


def test_office_evidence_hashes_and_geometry_are_pinned(root):
    data = json.loads((root / "docs/calibration/font-collections/evidence.json").read_text())
    fixture = root / data["fixture"]
    assert hashlib.sha256(fixture.read_bytes()).hexdigest() == data["fixture_sha256"]
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
          "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
    assert len(data["cases"]) == 6
    for case in data["cases"]:
        assert hashlib.sha256((root / case["render"]).read_bytes()).hexdigest() == case["render_sha256"]
        slide = etree.fromstring(read_part(fixture, f"ppt/slides/slide{case['slide']}.xml").encode())
        shape, = slide.findall(".//p:sp", ns)
        assert shape.find("p:nvSpPr/p:cNvPr", ns).get("name") == case["shape"]
        extent = shape.find("p:spPr/a:xfrm/a:ext", ns)
        assert int(extent.get("cx")) / 12700 == case["width_pt"]
        assert int(extent.get("cy")) / 12700 == 60
        body = shape.find("p:txBody/a:bodyPr", ns)
        assert body.get("wrap") == "none" and body.find("a:noAutofit", ns) is not None
        assert all(body.get(name) == "0" for name in ("lIns", "rIns", "tIns", "bIns"))
        run = shape.find(".//a:rPr", ns)
        assert run.get("sz") == "2400"
        assert run.get("b") == str(int(case["bold"]))
        assert run.get("i") == str(int(case["italic"]))
        assert run.find("a:latin", ns).get("typeface") == "Times"
        assert "".join(shape.xpath(".//a:t/text()", namespaces=ns)) == case["text"]


def test_office_geometry_with_synthetic_collection_on_every_platform(root, tmp_path, monkeypatch):
    # Original empty glyphs, approximate advance ratios, no proprietary font
    # data or installed Office needed in CI. The separate PNGs are Office proof.
    coll = TTCollection()
    for style, r, w in [("Regular", 333, 944), ("Bold", 444, 1000),
                        ("Italic", 389, 833), ("Bold Italic", 389, 889)]:
        f = make_font("Synthetic Times Test", style, 500)
        f["hmtx"].metrics["g114"] = (r, 0)
        f["hmtx"].metrics["g87"] = (w, 0)
        coll.fonts.append(f)
    coll.save(tmp_path / "office-ratios.ttc")
    coll.close()
    monkeypatch.setattr(fonts, "FONT_DIRS", (str(tmp_path),))
    monkeypatch.setattr(fonts, "_fc_match", lambda pattern: None)
    path = root / "corpus/pptx-font-collections.pptx"
    edits = {f"ppt/slides/slide{i}.xml": read_part(path, f"ppt/slides/slide{i}.xml")
             .replace('typeface="Times"', 'typeface="Synthetic Times Test"').encode()
             for i in range(1, 7)}
    changed = repack(path, tmp_path / "synthetic.pptx", edits)
    findings = [f for f in check_pptx(changed) if f.severity >= Severity.WARN]
    assert [(f.code, f.severity, f.where) for f in findings] == [
        ("PPT003", Severity.ERROR, "slide2/OVER_bold"),
        ("PPT003", Severity.ERROR, "slide5/OVER_bold_italic"),
    ]
