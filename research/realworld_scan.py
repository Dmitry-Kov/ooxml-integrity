"""Run the checker over public real-world test corpora.

The labelled corpus in evidence/ is synthetic. This script measures noise and
detection on documents other tools' projects collected from real producers.

    python research/realworld_scan.py fetch corpora/          # pinned sparse clones
    python research/realworld_scan.py scan corpora/ --json out.json
    python research/realworld_scan.py noop corpora/            # python-docx open/save
    python research/realworld_scan.py mutate corpora/          # destructive vs careful

The corpora are referenced by commit and path, not vendored: their licences
differ from this repository's. `noop` and `mutate` need python-docx (dev extra).
Results for the pinned commits are recorded in docs/real-world-corpora.md.
"""
from __future__ import annotations

import argparse
import collections
import json
import os
import re
import subprocess
import sys
import tempfile
import time
import warnings
import zipfile
from pathlib import Path

from ooxml_integrity import check, check_pptx, compare

SOURCES = [
    ("libreoffice", "https://github.com/LibreOffice/core.git",
     "382bad999f5d43448458ce2a3170a24c1781c867",
     ["sw/qa/extras/ooxmlexport/data", "sw/qa/extras/ooxmlimport/data",
      "sd/qa/unit/data/pptx"]),
    ("apache-poi", "https://github.com/apache/poi.git",
     "3b299acbf0195299f4cca3ba4207dd36b66d3210",
     ["test-data/document", "test-data/slideshow"]),
    ("docx4j", "https://github.com/plutext/docx4j.git",
     "49e66dd85bad4d2292b42222bf4867adea637a3a",
     ["docx4j-samples-docx4j/sample-docs", "docx4j-core-tests/src/test/resources",
      "docx4j-samples-docx-export-fo/sample-docs", "docs"]),
    ("open-xml-sdk", "https://github.com/dotnet/Open-XML-SDK.git",
     "431ab05cf160248cc3885a4a766026d4f8243792",
     ["test/DocumentFormat.OpenXml.Tests.Assets/assets/TestFiles"]),
    ("python-docx", "https://github.com/python-openxml/python-docx.git",
     "e45454602b53e8e572b179ccf1c91093ec9f4ed7",
     ["features/steps/test_files", "tests/test_files"]),
    ("python-pptx", "https://github.com/scanny/python-pptx.git",
     "278b47b1dedd5b46ee84c286e77cdfb0bf4594be",
     ["features/steps/test_files", "tests/test_files"]),
]
DECKS = (".pptx", ".potx", ".ppsx")


def fetch(root: Path) -> None:
    root.mkdir(parents=True, exist_ok=True)
    for name, url, commit, paths in SOURCES:
        target = root / name
        if (target / ".git").exists():
            print(f"{name}: present")
            continue
        run = lambda *a: subprocess.run(["git", "-C", str(target), *a], check=True)
        target.mkdir()
        run("init", "-q")
        run("remote", "add", "origin", url)
        run("sparse-checkout", "set", *paths)
        run("fetch", "-q", "--depth", "1", "--filter=blob:none", "origin", commit)
        run("checkout", "-q", "FETCH_HEAD")
        print(f"{name}: {commit[:12]}")


def documents(root: Path, suffixes=(".docx",) + DECKS) -> list[Path]:
    found = set()
    for directory, _, names in os.walk(root, followlinks=True):
        if "/.git" in directory:
            continue
        for name in names:
            if name.lower().endswith(suffixes):
                found.add(Path(directory, name).resolve())
    return sorted(found)


def producer(path: Path) -> str:
    try:
        app = zipfile.ZipFile(path).read("docProps/app.xml").decode("utf-8", "replace")
    except Exception:
        return "unknown"
    name = re.search(r"<Application>([^<]*)</Application>", app)
    ver = re.search(r"<AppVersion>([^<]*)</AppVersion>", app)
    label = (name.group(1) if name else "unknown").split("$")[0].strip()
    return f"{label} {ver.group(1)}" if ver else label


def scan(root: Path, output: Path | None, examples: int) -> None:
    rows = []
    rules = collections.Counter()
    files = collections.defaultdict(set)
    producers = collections.defaultdict(collections.Counter)
    samples = collections.defaultdict(list)
    started = time.time()
    for path in documents(root):
        try:
            findings = check_pptx(path) if path.suffix.lower() in DECKS else check(path)
        except Exception as error:  # a crash is a finding about the checker
            rows.append({"path": str(path), "crash": repr(error)})
            continue
        prod = producer(path)
        kept = [f for f in findings if f.severity.value != "info"]
        for f in kept:
            key = (path.suffix.lower(), f.code, f.severity.value)
            rules[key] += 1
            files[key].add(path)
            producers[key][prod] += 1
            if len(samples[key]) < examples:
                samples[key].append(f"{path.name}: {f.message[:110]}")
        rows.append({"path": str(path), "producer": prod,
                     "findings": [f.as_dict() for f in findings]})
    for suffix in (".docx", ".pptx"):
        mine = [r for r in rows if r["path"].lower().endswith(suffix)]
        failing = sum(1 for r in mine
                      if any(f["severity"] == "error" for f in r.get("findings", [])))
        print(f"{suffix}: {len(mine)} files, {failing} with errors")
    for key, count in sorted(rules.items(), key=lambda kv: -len(files[kv[0]])):
        top = ", ".join(f"{p} x{n}" for p, n in producers[key].most_common(3))
        print(f"\n{key[0]} {key[1]} {key[2]}: {count} in {len(files[key])} files [{top}]")
        for line in samples[key]:
            print(f"    {line}")
    crashes = [r for r in rows if "crash" in r]
    for row in crashes[:10]:
        print(f"CRASH {row['path']}: {row['crash']}")
    print(f"\n{time.time() - started:.0f}s")
    if output:
        output.write_text(json.dumps(rows, indent=1))


def _saved(document, directory: str, name: str) -> str:
    out = os.path.join(directory, name)
    document.save(out)
    return out


def noop(root: Path) -> None:
    import docx
    warnings.filterwarnings("ignore")
    compared = findings = 0
    failures = []
    with tempfile.TemporaryDirectory() as tmp:
        for i, path in enumerate(documents(root, (".docx",))):
            try:
                out = _saved(docx.Document(str(path)), tmp, f"{i}.docx")
            except Exception:
                continue  # python-docx cannot open it; not a checker result
            try:
                found = list(compare(path, out))
            except Exception as error:
                failures.append((path.name, repr(error)[:120]))
                continue
            compared += 1
            findings += len(found)
            for f in found:
                print(f"{path.name}: {f.code} {f.message[:100]}")
    print(f"{compared} python-docx open/save pairs compared, {findings} findings")
    for name, error in failures:
        print(f"comparison failed: {name}: {error}")


KINDS = {
    "comment": ".//w:commentRangeStart|.//w:commentReference",
    "footnote": ".//w:footnoteReference",
    # a paragraph-mark revision (in w:pPr) survives Paragraph.text; content counts
    "revision": ".//w:ins[not(ancestor::w:pPr)]|.//w:del[not(ancestor::w:pPr)]",
}


def mutate(root: Path) -> None:
    import docx
    warnings.filterwarnings("ignore")
    stats = collections.defaultdict(collections.Counter)
    missed, flagged = [], []
    with tempfile.TemporaryDirectory() as tmp:
        for i, path in enumerate(documents(root, (".docx",))):
            try:
                known = {(f.code, f.message) for f in check(path)}
                docx.Document(str(path))
            except Exception:
                continue
            for kind, query in KINDS.items():
                edited = docx.Document(str(path))
                targets = [p for p in edited.paragraphs
                           if p._p.xpath(query) and p.text.strip()]
                if not targets:
                    continue
                targets[0].text = targets[0].text + " (updated)"  # the agent edit
                out = _saved(edited, tmp, f"{i}-{kind}-bad.docx")
                new = [f for f in list(check(out)) + list(compare(path, out))
                       if f.severity.value == "error" and (f.code, f.message) not in known]
                stats[kind]["destructive"] += 1
                stats[kind]["detected"] += bool(new)
                if not new:
                    missed.append((kind, path.name))
                careful = docx.Document(str(path))
                run = next((r for p in careful.paragraphs for r in p.runs
                            if len(r.text.strip()) > 3), None)
                if run is None:
                    continue
                run.text = run.text + " (updated)"
                out = _saved(careful, tmp, f"{i}-{kind}-good.docx")
                new = [f for f in list(check(out)) + list(compare(path, out))
                       if f.severity.value in ("error", "warn")
                       and (f.code, f.message) not in known]
                stats[kind]["careful"] += 1
                stats[kind]["careful flagged"] += bool(new)
                if new:
                    flagged.append((kind, path.name, new[0].code))
    for kind, counts in stats.items():
        print(kind, dict(counts))
    print("missed:", missed)
    print("careful edits flagged:", flagged)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    sub = parser.add_subparsers(dest="command", required=True)
    for name in ("fetch", "scan", "noop", "mutate"):
        p = sub.add_parser(name)
        p.add_argument("root", type=Path)
        if name == "scan":
            p.add_argument("--json", type=Path)
            p.add_argument("--examples", type=int, default=3)
    args = parser.parse_args()
    if args.command == "fetch":
        fetch(args.root)
    elif args.command == "scan":
        scan(args.root, args.json, args.examples)
    elif args.command == "noop":
        noop(args.root)
    else:
        mutate(args.root)
    return 0


if __name__ == "__main__":
    sys.exit(main())
