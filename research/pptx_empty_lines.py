#!/usr/bin/env python3
"""Empty paragraphs and empty lines in .pptx text: build the probe deck, read a render.

The question: how tall is an empty line, and when does it move text that is
drawn? An empty paragraph has no runs, so its height can only come from its
mark (`a:endParaRPr`, else the inherited size); an empty line between two
breaks has the second break's size. Whether blank rows above the first line
of text or below the last one move any glyph depends on the anchor.

Each slide holds one text box with zero insets and no autofit, named for its
probe. Glyph lines start with distinct letters (`Abc` above, `Xyz` below), so
a PDF render gives their positions; an empty line is measured by how far it
pushes what follows.

    python research/pptx_empty_lines.py build corpus/pptx-empty-lines.pptx
    python research/pptx_empty_lines.py measure corpus/pptx-empty-lines.pptx \
        render.pdf --renderer "LibreOffice 25.8.4.2" --out measurements.json

`measure` writes, per probe, the top of the first `A` and the bottom of the
last `X` relative to the top of the box, as the renderer drew them.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from lxml import etree

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_pptx_corpus import _freeze_timestamps  # noqa: E402

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
FONT = '<a:latin typeface="Calibri"/>'


def run(text: str, sz: int = 18) -> str:
    return f'<a:r><a:rPr lang="en-US" sz="{sz * 100}">{FONT}</a:rPr><a:t>{text}</a:t></a:r>'


def br(sz: int = 18) -> str:
    return f'<a:br><a:rPr lang="en-US" sz="{sz * 100}">{FONT}</a:rPr></a:br>'


def para(body: str = "", end: int | None = 18, ppr: str = "") -> str:
    mark = f'<a:endParaRPr lang="en-US" sz="{end * 100}">{FONT}</a:endParaRPr>' if end else ""
    return f"<a:p>{ppr}{body}{mark}</a:p>"


def spacing(before: int = 0, after: int = 0) -> str:
    parts = []
    if before:
        parts.append(f'<a:spcBef><a:spcPts val="{before * 100}"/></a:spcBef>')
    if after:
        parts.append(f'<a:spcAft><a:spcPts val="{after * 100}"/></a:spcAft>')
    return f"<a:pPr>{''.join(parts)}</a:pPr>"


TOP, BOT = run("Abc"), run("Xyz")
EMPTY36 = para("", 36)

#: Height probes: top-anchored, in a box far taller than the text.
TALL = 468
MEASURE = {
    "base": [para(TOP), para(BOT)],
    "empty18": [para(TOP), para("", 18), para(BOT)],
    "empty36": [para(TOP), para("", 36), para(BOT)],
    "empty10": [para(TOP), para("", 10), para(BOT)],
    "empty_inherited": [para(TOP), para("", None), para(BOT)],
    "empty18_spacing": [para(TOP), para("", 18, spacing(12, 6)), para(BOT)],
    "two_empty36": [para(TOP), EMPTY36, EMPTY36, para(BOT)],
    "breaks_base": [para(TOP + br(36) + BOT)],
    "break_break36": [para(TOP + br(36) + br(36) + BOT)],
    "break_break10": [para(TOP + br(10) + br(10) + BOT)],
    "trailing_break_mark36": [para(TOP + br(18), 36), para(BOT)],
    "trailing_break_mark10": [para(TOP + br(18), 10), para(BOT)],
    "solo": [para(BOT)],
    "leading_empty36": [EMPTY36, para(BOT)],
    "mark36_on_text": [para(TOP, 36), para(BOT)],
}

#: Pairs whose difference in where `Xyz` is drawn is one probe's extra height.
DELTAS = [
    ("empty18", "base"), ("empty36", "base"), ("empty10", "base"),
    ("empty_inherited", "base"), ("empty18_spacing", "base"), ("two_empty36", "base"),
    ("break_break36", "breaks_base"), ("break_break10", "breaks_base"),
    ("trailing_break_mark36", "base"), ("trailing_break_mark10", "base"),
    ("leading_empty36", "solo"), ("mark36_on_text", "base"),
]

#: Anchor probes: (anchor, box height, paragraphs). Two 18pt lines of text are
#: 43.2pt; two 36pt empty paragraphs another 86.4pt.
ANCHORS: dict[str, tuple[str, int, list[str]]] = {}
for _anchor in ("t", "ctr", "b"):
    ANCHORS.update({
        f"{_anchor}_text": (_anchor, 100, [para(TOP), para(BOT)]),
        f"{_anchor}_trailing": (_anchor, 100, [para(TOP), para(BOT), EMPTY36, EMPTY36]),
        f"{_anchor}_leading": (_anchor, 100, [EMPTY36, EMPTY36, para(TOP), para(BOT)]),
        f"{_anchor}_trailing_break": (_anchor, 60, [para(TOP), para(BOT + br(36), 36)]),
        f"{_anchor}_first_space_before": (_anchor, 60, [para(TOP, ppr=spacing(36)), para(BOT)]),
        f"{_anchor}_last_space_after": (_anchor, 60, [para(TOP), para(BOT, ppr=spacing(0, 36))]),
    })

BOX_LEFT, BOX_WIDTH = 72, 576
BOX_TOP = {"measure": 36, "anchor": 200}


def probes() -> dict[str, dict]:
    out = {name: {"anchor": "t", "box_pt": TALL, "box_top_pt": BOX_TOP["measure"],
                  "paragraphs": paras} for name, paras in MEASURE.items()}
    out.update({name: {"anchor": anchor, "box_pt": h, "box_top_pt": BOX_TOP["anchor"],
                       "paragraphs": paras} for name, (anchor, h, paras) in ANCHORS.items()})
    return out


def build(out: Path) -> None:
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation()
    for name, probe in probes().items():
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Pt(BOX_LEFT), Pt(probe["box_top_pt"]),
                                       Pt(BOX_WIDTH), Pt(probe["box_pt"]))
        box.name = name
        old = box.text_frame._txBody
        xml = (f'<p:txBody xmlns:p="{P}" xmlns:a="{A}">'
               f'<a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" '
               f'anchor="{probe["anchor"]}"><a:noAutofit/></a:bodyPr><a:lstStyle/>'
               + "".join(probe["paragraphs"]) + "</p:txBody>")
        old.getparent().replace(old, etree.fromstring(xml))
    prs.save(out)
    _freeze_timestamps(str(out))


def measure(deck: Path, pdf: Path, renderer: str) -> dict:
    import pdfplumber

    from ooxml_integrity.pptx_layout import read_deck

    names = [s.name for s in sorted(read_deck(deck).shapes, key=lambda s: s.slide)]
    table = probes()
    rows = {}
    with pdfplumber.open(pdf) as doc:
        if len(doc.pages) != len(names):
            raise SystemExit(f"{pdf}: {len(doc.pages)} pages for {len(names)} slides")
        for name, page in zip(names, doc.pages):
            top = table[name]["box_top_pt"]
            a = [c["top"] for c in page.chars if c["text"] == "A"]
            x = [c["bottom"] for c in page.chars if c["text"] == "X"]
            rows[name] = {
                "a_top_pt": round(min(a) - top, 2) if a else None,
                "x_bottom_pt": round(max(x) - top, 2) if x else None,
            }
    return {"renderer": renderer, "probes": rows}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    sub = parser.add_subparsers(dest="command", required=True)
    b = sub.add_parser("build", help="write the probe deck")
    b.add_argument("out", type=Path)
    m = sub.add_parser("measure", help="read glyph positions from a PDF render of the deck")
    m.add_argument("deck", type=Path)
    m.add_argument("pdf", type=Path)
    m.add_argument("--renderer", required=True)
    m.add_argument("--out", type=Path)
    args = parser.parse_args(argv)
    if args.command == "build":
        build(args.out)
        print(args.out)
        return 0
    result = measure(args.deck, args.pdf, args.renderer)
    text = json.dumps(result, indent=1) + "\n"
    if args.out:
        args.out.write_text(text)
    else:
        sys.stdout.write(text)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
