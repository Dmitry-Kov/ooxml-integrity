#!/usr/bin/env python3
"""Frozen, fixture-specific boundary benchmark; no production checker changes."""
from __future__ import annotations

import argparse
from copy import deepcopy
import importlib.metadata as metadata
import io
import json
import platform
import subprocess
import sys
import traceback
from contextlib import redirect_stdout, redirect_stderr
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from lxml import etree as E
from research import docx_benchmark as b
from research import revision_evidence as r

BASE = ROOT / "evidence/docx-benchmark-boundaries"
PROTOCOL = BASE / "protocol.json"
W, NS, MAIN = r.W, r.NS, r.MAIN
TASKS = {
    "save": {"old": None, "new": None},
    "split": {"old": "EDITBEFORE", "new": "DONEAFTER", "adeu_new": "**DONE**_AFTER_"},
    "comment": {"old": "Review the delivery terms.", "new": "Review the payment terms."},
    "table": {"old": "TABLEBEFORE", "new": "TABLEAFTER"},
}
ADAPTERS = ("adeu-sdk-boundaries-v1", "python-docx-targets-v2")


def target(root, task):
    if task == "split":
        return root.findall("w:body/w:p", NS)[2]
    if task == "comment":
        return root.findall("w:body/w:p", NS)[5]
    if task == "table":
        return root.findall("w:body/w:tbl/w:tr", NS)[1].findall("w:tc", NS)[1].find("w:p", NS)
    raise ValueError(task)


def derived_parts(profile):
    parts = r.read(r.BASE / "sources" / f"{profile}.docx")
    root = E.fromstring(parts[MAIN])
    paragraph = target(root, "split")
    original = paragraph[0]
    prefix, suffix = original.find(W + "t").text.split("EDITBEFORE")
    paragraph.remove(original)
    for i, (text, style) in enumerate(((prefix, None), ("EDIT", "b"), ("BEFORE", "i"), (suffix, None))):
        run = E.Element(W + "r")
        text_node = E.SubElement(run, W + "t")
        text_node.set(b.SPACE, "preserve")
        text_node.text = text
        if style:
            props = E.Element(W + "rPr")
            props.append(E.Element(W + style))
            run.insert(0, props)
        paragraph.insert(i, run)
    target(root, "table").findall("w:r/w:t", NS)[-1].text = "TABLEBEFORE"
    parts[MAIN] = r.xml(root)
    return parts


def freeze():
    if PROTOCOL.exists():
        raise ValueError("Already frozen")
    sources = {}
    for profile in r.PROFILES:
        original = r.BASE / "sources" / f"{profile}.docx"
        path = BASE / "sources" / f"{profile}.docx"
        if path.exists():
            raise ValueError("Existing derived source")
        r.write(derived_parts(profile), path)
        sources[profile] = {"path": str(path.relative_to(ROOT)), "sha256": b.digest(path),
                            "original": str(original.relative_to(ROOT)), "original_sha256": b.digest(original)}
    frozen = [BASE / "PROTOCOL.md", BASE / "prompts.json", Path(__file__),
              ROOT / "research/capture_docx_boundaries_agent.py", Path(b.__file__), Path(r.__file__),
              ROOT / "research/capture_docx_agent.py"]
    b.save_json(PROTOCOL, {"schema": 1, "declared_at": b.utc(), "tasks": TASKS, "sources": sources,
                          "base_commit": subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip(),
                          "frozen_files": {str(p.relative_to(ROOT)): b.digest(p) for p in frozen},
                          "checker_source_sha256": b.tree_hash(ROOT / "src/ooxml_integrity")})


def protocol():
    data = json.loads(PROTOCOL.read_text())
    for name, digest in data["frozen_files"].items():
        if b.digest(ROOT / name) != digest:
            raise ValueError(f"Frozen file drift: {name}")
    for s in data["sources"].values():
        if b.digest(ROOT / s["path"]) != s["sha256"] or b.digest(ROOT / s["original"]) != s["original_sha256"]:
            raise ValueError("Source drift")
    if b.tree_hash(ROOT / "src/ooxml_integrity") != data["checker_source_sha256"]:
        raise ValueError("Checker drift")
    return data


def python_edit(source, output, task):
    from docx import Document
    doc = Document(source)
    if task != "save":
        p = doc.tables[0].rows[1].cells[1].paragraphs[0] if task == "table" else doc.paragraphs[2 if task == "split" else 5]
        if task == "split":
            runs = p.runs
            matches = [i for i in range(len(runs) - 1) if runs[i].text == "EDIT" and runs[i + 1].text == "BEFORE"]
            if len(matches) != 1:
                raise b.Unsupported("Expected adjacent EDIT/BEFORE ordinary runs")
            runs[matches[0]].text = "DONE"
            runs[matches[0] + 1].text = "AFTER"
        else:
            old, new = TASKS[task]["old"], TASKS[task]["new"]
            matches = [run for run in p.runs if old in run.text]
            if len(matches) != 1 or matches[0].text.count(old) != 1:
                raise b.Unsupported("Expected one ordinary run at the declared location")
            matches[0].text = matches[0].text.replace(old, new, 1)
    doc.save(output)
    return {"operation": "Document/paragraphs or tables.rows.cells.paragraphs/Run.text/save", "task": task}


def adeu_edit(source, output, task):
    from adeu import ModifyText, RedlineEngine
    engine = RedlineEngine(io.BytesIO(source.read_bytes()), author=b.AUTHOR)
    stats, parameters = None, None
    if task != "save":
        spec = TASKS[task]
        change = ModifyText(target_text=spec["old"], new_text=spec.get("adeu_new", spec["new"]),
                            match_mode="strict", regex=False, comment=None)
        parameters = change.model_dump()
        try:
            stats = engine.process_batch([change], partial=False)
        except Exception as exc:
            if type(exc).__name__ == "BatchValidationError":
                raise b.Rejected(str(exc)) from exc
            raise
        if stats.get("failed") or stats.get("edits_applied", 0) != 1:
            raise b.Rejected(json.dumps(stats, default=str))
    output.write_bytes(engine.save_to_stream().getvalue())
    return {"stats": stats, "change": parameters, "author": b.AUTHOR, "partial": False,
            "engine_options": "3.0.4 defaults except author"}


def capture(adapter, directory):
    data = protocol()
    pins = {"python-docx": "1.2.0", "lxml": "6.1.3"}
    if adapter == ADAPTERS[0]:
        pins["adeu"] = "3.0.4"
    if any(metadata.version(k) != v for k, v in pins.items()):
        raise ValueError("Dependency drift")
    directory.mkdir(parents=True, exist_ok=False)
    (directory / "capture-harness.py").write_bytes(Path(__file__).read_bytes())
    import docx
    implementations = {"docx": b.tree_hash(Path(docx.__file__).parent)}
    if adapter == ADAPTERS[0]:
        import adeu
        implementations["adeu"] = b.tree_hash(Path(adeu.__file__).parent)
    receipt = {"adapter": adapter, "started_at": b.utc(), "protocol_sha256": b.digest(PROTOCOL),
               "harness_sha256": b.digest(__file__), "command": [sys.executable, *sys.argv],
               "cwd": str(ROOT), "python": platform.python_version(), "platform": platform.platform(),
               "distributions": dict(sorted((d.metadata["Name"], d.version) for d in metadata.distributions())),
               "implementation_sha256": implementations, "cases": []}
    b.save_json(directory / "started.json", receipt)
    for profile, source in data["sources"].items():
        for task in TASKS:
            ident = f"{profile}-{task}"
            output = directory / f"{ident}.docx"
            record = {"id": ident, "profile": profile, "task": task, "source": source["path"],
                      "source_sha256": source["sha256"], "started_at": b.utc()}
            stdout, stderr = io.StringIO(), io.StringIO()
            try:
                with redirect_stdout(stdout), redirect_stderr(stderr):
                    record["tool_result"] = (adeu_edit if adapter == ADAPTERS[0] else python_edit)(ROOT / source["path"], output, task)
                if not output.is_file():
                    raise RuntimeError("Missing output")
                record["status"] = "ok"
            except (b.Unsupported, b.Rejected) as exc:
                record.update(status="unsupported" if isinstance(exc, b.Unsupported) else "rejected", error=str(exc))
            except Exception as exc:
                record.update(status="error", error=f"{type(exc).__name__}: {exc}")
                traceback.print_exc(file=stderr)
            record.update(finished_at=b.utc(), stdout=stdout.getvalue(), stderr=stderr.getvalue(),
                          output=output.name if output.exists() else None,
                          output_sha256=b.digest(output) if output.exists() else None)
            b.save_json(directory / f"{ident}.json", record)
            receipt["cases"].append(record)
            print(adapter, ident, record["status"], flush=True)
    receipt["finished_at"] = b.utc()
    b.save_json(directory / "capture.json", receipt)


def expected_parts(parts, task):
    """Independent fixture expectation; no editor or checker imports."""
    parts = dict(parts)
    root = E.fromstring(parts[MAIN])
    p = target(root, task)
    for t in p.findall("w:r/w:t", NS):
        if task == "split" and t.text in ("EDIT", "BEFORE"):
            t.text = {"EDIT": "DONE", "BEFORE": "AFTER"}[t.text]
        elif task != "split":
            t.text = (t.text or "").replace(TASKS[task]["old"], TASKS[task]["new"])
    parts[MAIN] = r.xml(root)
    return parts


def plain_tokens(run):
    if run.tag != W + "r" or not all(n.tag in (W + "rPr", W + "t") for n in run):
        return None
    props = run.find(W + "rPr")
    if len(run.findall(W + "rPr")) > 1:
        return None
    if props is not None:
        props = deepcopy(props)
        for prop in props:
            if prop.tag in (W + "b", W + "i") and prop.get(W + "val") in ("1", "true", "on"):
                prop.attrib.pop(W + "val")
        if len(props) == 0 and not props.attrib and not (props.text or "").strip():
            props = None
    key = (tuple(sorted(run.attrib.items())), b.node_signature(props) if props is not None else None)
    tokens = []
    for t in run.findall(W + "t"):
        if len(t):
            return None
        attrs = tuple(sorted((k, v) for k, v in t.attrib.items() if k != b.SPACE))
        tokens.extend(("char", key, attrs, char) for char in (t.text or ""))
    return tokens


def paragraph_signature(p):
    tokens = []
    for node in p:
        text = plain_tokens(node)
        tokens.extend(text if text is not None else [("node", b.node_signature(node))])
    return (p.tag, tuple(sorted(p.attrib.items())), (p.text or "").strip(),
            (p.tail or "").strip(), tokens)


def package(parts, task):
    parts = dict(parts)
    if task != "save":
        root = E.fromstring(parts[MAIN])
        p = target(root, task)
        sentinel = E.Element("benchmark-target-signature")
        sentinel.text = b.json_bytes(paragraph_signature(p)).decode()
        p.getparent().replace(p, sentinel)
        parts[MAIN] = r.xml(root)
    return b.package_facts(parts, "save")


def project_new_revisions(before, after, task):
    """Permit only declared direct new wrappers; keep all other markup exact."""
    if task == "save":
        return after, after, 0, []
    source = E.fromstring(before[MAIN])
    old_ids = {n.get(W + "id") for n in source.iter() if n.tag in (W + "ins", W + "del")}
    views = []
    counts, issues = 0, []
    for current in (True, False):
        root = E.fromstring(after[MAIN])
        p = target(root, task)
        count = 0
        for node in list(p):
            if node.tag not in (W + "ins", W + "del") or node.get(W + "id") in old_ids:
                continue
            attrs = {W + n for n in ("id", "author", "date")}
            extension = "{http://schemas.microsoft.com/office/word/2023/wordml/word16du}dateUtc"
            if node.get(extension) == node.get(W + "date"):
                attrs.add(extension)
            runs = deepcopy(node)
            for t in runs.iter(W + "delText"):
                t.tag = W + "t"
            valid = (node.get(W + "id") is not None and node.get(W + "author") == b.AUTHOR
                     and not set(node.attrib) - attrs and len(node) > 0
                     and all(plain_tokens(child) is not None for child in runs))
            if not valid:
                if current:
                    issues.append("unrecognized_new_revision")
                continue
            count += 1
            if (node.tag == W + "ins") == current:
                for t in node.iter(W + "delText"):
                    t.tag = W + "t"
                r.unwrap(node)
            else:
                p.remove(node)
        projected = dict(after)
        projected[MAIN] = r.xml(root)
        views.append(projected)
        counts = count
    return views[0], views[1], counts, issues


def current_text(parts, task):
    root = E.fromstring(parts[MAIN])
    p = target(root, task)
    return "".join(t.text or "" for t in p.iter(W + "t")
                   if not any(a.tag == W + "del" for a in t.iterancestors()))


def assess(source, output, task, adapter):
    before, after = r.read(source), r.read(output)
    current, original, new_count, issues = project_new_revisions(before, after, task)
    raw_facts = r.facts(after)
    if task == "save":
        completed = r.facts(before)["text"] == raw_facts["text"]
        expected = before
    else:
        observed = current_text(after, task)
        completed = observed.count(TASKS[task]["new"]) == 1 and observed.count(TASKS[task]["old"]) == 0
        # An unchanged task target is not collateral loss. Never accept it as completion.
        unchanged_target = observed.count(TASKS[task]["old"]) == 1 and TASKS[task]["new"] not in observed
        expected = before if unchanged_target else expected_parts(before, task)
    case = {"id": task, "action": "edit", "cohort": "bounded_current_view", "allowed_revision_losses": [],
            "allowed_text_changes": []}
    oracle = r.oracle(expected, current, case)
    old_package, old_meta = package(expected, task)
    new_package, new_meta = package(current, task)
    changed = [p for p in sorted(set(old_package) | set(new_package)) if old_package.get(p) != new_package.get(p)]
    old_view_changes = []
    if new_count:
        a, _ = package(before, task)
        z, _ = package(original, task)
        old_view_changes = [p for p in sorted(set(a) | set(z)) if a.get(p) != z.get(p)]
    tracked_missing = task != "save" and adapter == ADAPTERS[0] and not new_count
    violations = {**oracle["violations"], "protected_package_changes": changed,
                  "original_revision_view_changes": old_view_changes, "new_revision_shape": issues,
                  "raw_invalid_revision_text": raw_facts["invalid"], "raw_duplicate_revision_ids": raw_facts["duplicates"]}
    return {"requested_change_completed": bool(completed and not tracked_missing),
            "protected_content_preserved": not any(violations.values()), "new_revision_wrappers": new_count,
            "tracked_edit_missing": tracked_missing, "violations": violations, "revision_oracle": oracle,
            "normalized_package_sha256": r.digest(b.json_bytes(new_package)),
            "changed_raw_parts": [p for p in sorted(set(before) | set(after)) if before.get(p) != after.get(p)],
            "mutable_metadata_changed": old_meta != new_meta,
            "raw_part_sha256": {p: r.digest(blob) for p, blob in sorted(after.items())}}


def inspect(source, output, task, adapter):
    result = assess(source, output, task, adapter)
    findings = b.checker_findings(source, output)
    actionable = [f for f in findings if f["severity"] in ("error", "warn")]
    return {**result, "checker_findings": findings, "actionable_findings": actionable,
            "checker_detection": "no_independent_defect" if result["protected_content_preserved"] else
            "findings_require_mapping" if actionable else "checker_silent_on_undeclared_change"}


def evaluate(directory):
    data = protocol()
    receipt = json.loads((directory / "capture.json").read_text())
    if receipt["protocol_sha256"] != b.digest(PROTOCOL) or receipt["harness_sha256"] != b.digest(directory / "capture-harness.py"):
        raise ValueError("Capture declaration drift")
    expected = {f"{p}-{t}" for p in data["sources"] for t in TASKS}
    if len(receipt["cases"]) != len(expected) or {c["id"] for c in receipt["cases"]} != expected:
        raise ValueError("Inventory drift")
    results = []
    for c in receipt["cases"]:
        if c != json.loads((directory / f"{c['id']}.json").read_text()):
            raise ValueError("Attempt drift")
        s = data["sources"][c["profile"]]
        if c["source"] != s["path"] or c["source_sha256"] != s["sha256"] or c["id"] != f"{c['profile']}-{c['task']}":
            raise ValueError("Task drift")
        result = {k: c[k] for k in ("id", "task", "profile", "status")}
        result.update(requested_change_completed=None, protected_content_preserved=None,
                      checker_detection="not_evaluated", overall_success=False)
        if c["output"] is not None:
            if c["output"] != f"{c['id']}.docx" or b.digest(directory / c["output"]) != c["output_sha256"]:
                raise ValueError("Output drift")
            try:
                result.update(inspect(ROOT / c["source"], directory / c["output"], c["task"], receipt["adapter"]))
                result["overall_success"] = c["status"] == "ok" and result["requested_change_completed"] and result["protected_content_preserved"]
            except (ValueError, KeyError, IndexError, E.XMLSyntaxError) as exc:
                result["evaluation_error"] = f"{type(exc).__name__}: {exc}"
        results.append(result)
    return {"adapter": receipt["adapter"], "capture_sha256": b.digest(directory / "capture.json"),
            "evaluator_sha256": b.digest(__file__), "checker_source_sha256": data["checker_source_sha256"],
            "source_findings": {p: b.checker_findings(ROOT / s["path"], ROOT / s["path"]) for p, s in data["sources"].items()},
            "cases": results}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    sub.add_parser("freeze")
    p = sub.add_parser("capture")
    p.add_argument("--adapter", choices=ADAPTERS, required=True)
    p.add_argument("--directory", type=Path, required=True)
    for name in ("evaluate", "verify"):
        p = sub.add_parser(name)
        p.add_argument("directory", type=Path)
    args = parser.parse_args()
    if args.command == "freeze":
        freeze()
    elif args.command == "capture":
        capture(args.adapter, args.directory.absolute())
    else:
        result = evaluate(args.directory)
        if args.command == "evaluate":
            b.save_json(args.directory / "evaluation.json", result)
        elif result != json.loads((args.directory / "evaluation.json").read_text()):
            raise ValueError("Evaluation drift")
        print(json.dumps({"attempts": len(result["cases"]), "successes": sum(c["overall_success"] for c in result["cases"])}))


if __name__ == "__main__":
    main()
