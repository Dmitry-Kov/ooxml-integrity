#!/usr/bin/env python3
"""Text boxes no taller than their insets: build the probe deck, read a render.

A shape whose height is no more than its top and bottom insets leaves no room
for text. Through 0.8.3 the checker reported nothing for it, because the
overflow ratio of a zero-height box was taken as 0. Each slide holds one text
box with one 18pt line (`AXbc`), no autofit, and the default insets (3.6pt top
and bottom, 7.2pt left and right) unless the probe says otherwise.

    python research/pptx_thin_boxes.py build corpus/pptx-thin-boxes.pptx
    python research/pptx_thin_boxes.py measure corpus/pptx-thin-boxes.pptx \
        render.pdf --renderer "LibreOffice 25.8.4.2" --out measurements.json

`measure` writes, per probe, the top of the `A` and the bottom of the `X`
relative to the top of the box, as the renderer drew them.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from lxml import etree

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_pptx_corpus import _freeze_timestamps  # noqa: E402
from pptx_empty_lines import A, P, para, run  # noqa: E402

LINE = para(run("AXbc"))
DEFAULT_INSETS = ""
NO_INSETS = ' lIns="0" tIns="0" rIns="0" bIns="0"'

#: name -> (anchor, box height in points, bodyPr inset attributes)
PROBES: dict[str, tuple[str, float, str]] = {
    "t_0": ("t", 0, DEFAULT_INSETS),
    "t_4": ("t", 4, DEFAULT_INSETS),
    "t_7.2": ("t", 7.2, DEFAULT_INSETS),
    "t_8": ("t", 8, DEFAULT_INSETS),
    "ctr_4": ("ctr", 4, DEFAULT_INSETS),
    "b_4": ("b", 4, DEFAULT_INSETS),
    "t_4_no_insets": ("t", 4, NO_INSETS),
    "t_40": ("t", 40, DEFAULT_INSETS),
}
BOX_LEFT, BOX_TOP, BOX_WIDTH = 72, 200, 576


def build(out: Path) -> None:
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation()
    for name, (anchor, height, insets) in PROBES.items():
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Pt(BOX_LEFT), Pt(BOX_TOP), Pt(BOX_WIDTH), Pt(height))
        box.name = name
        old = box.text_frame._txBody
        xml = (f'<p:txBody xmlns:p="{P}" xmlns:a="{A}">'
               f'<a:bodyPr wrap="square"{insets} anchor="{anchor}"><a:noAutofit/></a:bodyPr>'
               f"<a:lstStyle/>{LINE}</p:txBody>")
        old.getparent().replace(old, etree.fromstring(xml))
    prs.save(out)
    _freeze_timestamps(str(out))


def measure(deck: Path, pdf: Path, renderer: str) -> dict:
    import pdfplumber

    from ooxml_integrity.pptx_layout import read_deck

    names = [s.name for s in sorted(read_deck(deck).shapes, key=lambda s: s.slide)]
    rows = {}
    with pdfplumber.open(pdf) as doc:
        if len(doc.pages) != len(names):
            raise SystemExit(f"{pdf}: {len(doc.pages)} pages for {len(names)} slides")
        for name, page in zip(names, doc.pages):
            a = [c["top"] for c in page.chars if c["text"] == "A"]
            x = [c["bottom"] for c in page.chars if c["text"] == "X"]
            rows[name] = {"a_top_pt": round(min(a) - BOX_TOP, 2),
                          "x_bottom_pt": round(max(x) - BOX_TOP, 2)}
    return {"renderer": renderer, "probes": rows}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    sub = parser.add_subparsers(dest="command", required=True)
    b = sub.add_parser("build", help="write the probe deck")
    b.add_argument("out", type=Path)
    m = sub.add_parser("measure", help="read glyph positions from a PDF render of the deck")
    m.add_argument("deck", type=Path)
    m.add_argument("pdf", type=Path)
    m.add_argument("--renderer", required=True)
    m.add_argument("--out", type=Path)
    args = parser.parse_args(argv)
    if args.command == "build":
        build(args.out)
        print(args.out)
        return 0
    text = json.dumps(measure(args.deck, args.pdf, args.renderer), indent=1) + "\n"
    if args.out:
        args.out.write_text(text)
    else:
        sys.stdout.write(text)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
