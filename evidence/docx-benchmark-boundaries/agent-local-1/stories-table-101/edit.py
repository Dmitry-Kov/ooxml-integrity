from docx import Document

def replace_marker(input_path, output_path):
    doc = Document(input_path)
    
    tables = doc.tables
    if len(tables) < 1:
        raise Exception("No tables found in the document.")
    
    table = tables[0]
    if len(table.rows) < 2 or len(table.rows[1].cells) < 2:
        raise Exception("The specified cell does not exist.")
    
    cell = table.cell(1, 1)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            if run.text == "TABLEBEFORE":
                run.text = "TABLEAFTER"
                run._element.tag = 'w:t'
                run._element.set('xml:space', 'preserve')
    
    doc.save(output_path)

replace_marker("input.docx", "output.docx")
